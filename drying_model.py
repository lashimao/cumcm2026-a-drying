"""Conservative radial drying model. Inputs are read-only; SI units internally."""
from dataclasses import dataclass, asdict
from pathlib import Path
import hashlib
import math
import time
import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.sparse import lil_matrix, diags, bmat, csr_matrix


def numeric_sheet(filename):
    wb = load_workbook(filename, read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    values = np.asarray(rows[1:], dtype=float)
    if not np.isfinite(values).all():
        raise ValueError(f'Missing/nonfinite input in {Path(filename).name}')
    if not np.all(np.diff(values[:, 0]) > 0):
        raise ValueError('Input times must be strictly increasing')
    return rows[0], values


class Inputs:
    def __init__(self, directory):
        self.directory = Path(directory).resolve()
        _, a = numeric_sheet(self.directory / '附件1.xlsx')
        _, r = numeric_sheet(self.directory / '附件2.xlsx')
        self.ambient = a
        self.radius_data = r
        if a.shape != (241, 3) or r.shape != (145, 2):
            raise ValueError('Unexpected official input dimensions')
        if not (a[0, 0] == 0 and a[-1, 0] == 14400):
            raise ValueError('Unexpected ambient observation window')
        if not (r[0, 1] == 2 and np.all(np.diff(r[:, 1]) <= 0)):
            raise ValueError('Radius must start at 2 cm and not increase')
        self.radius_interpolator = PchipInterpolator(r[:, 0], r[:, 1] * .01,
                                                     extrapolate=False)
        tail = a[a[:, 0] >= 10800, 1:]
        self.tail = {
            'mean1h': tail.mean(axis=0),
            'last': a[-1, 1:],
            'mean2h': a[a[:, 0] >= 7200, 1:].mean(axis=0),
            'fast': np.array([tail[:, 0].max(), tail[:, 1].min()]),
            'slow': np.array([tail[:, 0].min(), tail[:, 1].max()]),
        }
        self.hashes = {f.name: hashlib.sha256(f.read_bytes()).hexdigest()
                       for f in sorted(self.directory.glob('*.xlsx'))}

    def boundary(self, t, mode='mean1h'):
        if t > self.ambient[-1, 0]:
            return self.tail[mode]
        return np.array([np.interp(t, self.ambient[:, 0], self.ambient[:, k])
                         for k in (1, 2)])

    def radius(self, t, moving):
        t = np.asarray(t, dtype=float)
        if not moving:
            return np.full_like(t, .02)
        return self.radius_interpolator(np.clip(t, 0, self.radius_data[-1, 0]))


@dataclass(frozen=True)
class Config:
    law: str = 'q23'
    moving: bool = False
    n: int = 160
    boundary: str = 'mean1h'
    d_scale: float = 1.0
    hm_scale: float = 1.0
    h_scale: float = 1.0
    rtol: float = 1e-7
    atol_t: float = 1e-7
    atol_c: float = 1e-9
    early_max_step: float = 60.0
    late_max_step: float = 300.0


def properties(c, temperature, law):
    c = np.asarray(c)
    if np.any(c <= 0) or np.any(np.asarray(temperature) <= -273.15):
        raise FloatingPointError('Coefficient evaluated outside physical domain')
    if law == 'q1':
        rho = np.full_like(c, 820.)
        cp = np.full_like(c, 2600.)
        k = np.full_like(c, .36)
        d = 7e-9 * np.exp(-.89 / c)
    elif law == 'q23':
        rho = 650 + 128 * c
        cp = 1450 + 2736 * c / (1 + c)
        k = .21 + .38 * c / (1 + c)
        d = 2.4e-3 * np.exp(-.45 / c - 3850 / (temperature + 273.15))
    elif law == 'q4':
        rho = 760 + 90 * c
        cp = 1850 + 2150 * c / (1 + c)
        k = .12 + .20 * c / (1 + c)
        d = 4.2e-4 * np.exp(-.30 / c - 3850 / (temperature + 273.15))
    else:
        raise ValueError(f'Unknown constitutive law {law}')
    return rho, cp, k, d


class RadialModel:
    def __init__(self, inputs, config, constant_d=None):
        self.inputs, self.config = inputs, config
        self.np = config.n + 1
        self.s = np.linspace(0, 1, self.np)
        faces = np.r_[0, (self.s[:-1] + self.s[1:]) / 2, 1]
        self.weights = np.diff(faces ** 2) / 2
        self.faces = faces
        self.constant_d = constant_d
        self.h = 25 * config.h_scale
        self.hm = 8e-7 * config.hm_scale
        self.sparsity = self.jacobian_pattern()

    def jacobian_pattern(self):
        m = self.np
        a = lil_matrix((2*m + 1, 2*m + 1), dtype=int)
        for j in range(m):
            neighbors = range(max(0, j-1), min(m, j+2))
            for k in neighbors:
                a[j, k] = a[j, m+k] = 1
                a[m+j, k] = a[m+j, m+k] = 1
        a[-1, 2*m-1] = 1
        return a.tocsr()

    def flux_divergence(self, u, coefficient, radius, external, exchange):
        harmonic = 2 * coefficient[:-1] * coefficient[1:] / (
            coefficient[:-1] + coefficient[1:])
        flux = np.empty(self.np + 1)
        flux[0] = 0
        flux[1:-1] = self.faces[1:-1] * harmonic * np.diff(u) * self.config.n
        flux[-1] = -radius * exchange * (u[-1] - external)
        return np.diff(flux) / (radius**2 * self.weights)

    def rhs(self, t, state):
        m = self.np
        temp, c = state[:m], state[m:2*m]
        rho, cp, k, d = properties(c, temp, self.config.law)
        if self.constant_d is not None:
            d = np.full_like(c, self.constant_d)
        d = d * self.config.d_scale
        radius = float(self.inputs.radius(t, self.config.moving))
        ta, ca = self.inputs.boundary(t, self.config.boundary)
        dt = self.flux_divergence(temp, k, radius, ta, self.h) / (rho * cp)
        dc = self.flux_divergence(c, d, radius, ca, self.hm)
        lost = 2 * self.hm / radius * (c[-1] - ca)
        return np.r_[dt, dc, lost]

    def jacobian(self, t, state):
        """Analytic sparse derivative; stable even after temperature equilibrates."""
        m, cfg = self.np, self.config
        temp, c = state[:m], state[m:2*m]
        rho, cp, k, d = properties(c, temp, cfg.law)
        if cfg.law == 'q1':
            drho = dcp = dk = np.zeros(m)
            alpha = .89
        elif cfg.law == 'q23':
            drho = np.full(m, 128.)
            dcp, dk, alpha = 2736/(1+c)**2, .38/(1+c)**2, .45
        else:
            drho = np.full(m, 90.)
            dcp, dk, alpha = 2150/(1+c)**2, .20/(1+c)**2, .30
        d *= cfg.d_scale
        dc_d = d*alpha/c**2
        dt_d = np.zeros(m) if cfg.law == 'q1' else d*3850/(temp+273.15)**2
        if self.constant_d is not None:
            d = np.full(m, self.constant_d*cfg.d_scale)
            dc_d = dt_d = np.zeros(m)
        radius = float(self.inputs.radius(t, cfg.moving))
        denom = radius**2*self.weights
        factor = self.faces[1:-1]*cfg.n
        def harmonic_parts(a):
            total = a[:-1]+a[1:]
            return 2*a[:-1]*a[1:]/total, 2*a[1:]**2/total**2, 2*a[:-1]**2/total**2
        def flux_matrix(left, right, boundary=0.):
            diagonal = np.zeros(m)
            diagonal[:-1] += left/denom[:-1]
            diagonal[1:] -= right/denom[1:]
            diagonal[-1] += boundary/denom[-1]
            return diags([-left/denom[1:], diagonal, right/denom[:-1]], [-1,0,1],
                         shape=(m,m), format='csr')
        kh, kl, kr = harmonic_parts(k)
        dh, dl, dr = harmonic_parts(d)
        delta_t, delta_c = np.diff(temp), np.diff(c)
        cap = rho*cp
        inv_cap = diags(1/cap)
        jtt = inv_cap @ flux_matrix(-factor*kh, factor*kh, -radius*self.h)
        jtc = inv_cap @ flux_matrix(factor*kl*dk[:-1]*delta_t,
                                    factor*kr*dk[1:]*delta_t)
        jtc -= diags(self.rhs(t, state)[:m]*(drho*cp+rho*dcp)/cap)
        jct = flux_matrix(factor*dl*dt_d[:-1]*delta_c,
                           factor*dr*dt_d[1:]*delta_c)
        jcc = flux_matrix(factor*(-dh+dl*dc_d[:-1]*delta_c),
                           factor*(dh+dr*dc_d[1:]*delta_c), -radius*self.hm)
        loss_c = csr_matrix(([2*self.hm/radius], ([0], [m-1])), shape=(1,m))
        return bmat([[jtt,jtc,csr_matrix((m,1))],
                     [jct,jcc,csr_matrix((m,1))],
                     [csr_matrix((1,m)),loss_c,csr_matrix((1,1))]], format='csc')

    def initial_state(self):
        return np.r_[np.full(self.np, 28.), np.full(self.np, 2.55), 0.]


class Trajectory:
    def __init__(self, model, segments, event_time, operational_time, elapsed):
        self.model, self.segments = model, segments
        self.event_time = event_time
        self.end_time = operational_time
        self.elapsed = elapsed

    def evaluate(self, times):
        scalar = np.ndim(times) == 0
        times = np.atleast_1d(times).astype(float)
        if np.min(times) < -1e-8 or np.max(times) > self.end_time + 1e-7:
            raise ValueError('Requested sampling outside solved interval')
        values = np.empty((2*self.model.np+1, len(times)))
        assigned = np.zeros(len(times), dtype=bool)
        for segment in self.segments:
            use = ((times >= segment.t[0] - 1e-8)
                   & (times <= segment.t[-1] + 1e-8) & ~assigned)
            if use.any():
                values[:, use] = segment.sol(times[use])
                assigned[use] = True
        if not assigned.all():
            raise RuntimeError('Unassigned trajectory sample')
        return values[:, 0] if scalar else values

    def sample(self, times, radii_cm):
        times = np.atleast_1d(times).astype(float)
        state = self.evaluate(times)
        m = self.model.np
        radius = self.model.inputs.radius(times, self.model.config.moving)
        grid = np.asarray(radii_cm) * .01
        temp, concentration = [], []
        for j, rad in enumerate(radius):
            query = grid / rad
            temp.append(np.interp(query, self.model.s, state[:m, j], right=np.nan))
            concentration.append(np.interp(query, self.model.s,
                                           state[m:2*m, j], right=np.nan))
        return np.asarray(temp), np.asarray(concentration)

    def diagnostics(self):
        times = np.unique(np.r_[np.linspace(0, self.end_time, 1201),
                                 self.end_time])
        state = self.evaluate(times)
        m, w = self.model.np, self.model.weights
        temp, c = state[:m], state[m:2*m]
        invariant = 2*w @ c + state[-1] - 2.55
        residuals = []
        for j in np.linspace(0, len(times)-1, 21, dtype=int):
            derivative = self.model.rhs(times[j], state[:, j])
            residuals.append(abs(2*w @ derivative[m:2*m] + derivative[-1]))
        end = self.evaluate(self.end_time)[m:2*m]
        return {
            'config': asdict(self.model.config),
            'threshold_seconds': self.event_time,
            'threshold_hours': None if self.event_time is None else self.event_time/3600,
            'operational_seconds': self.end_time,
            'end_c_max_unrounded': float(end.max()),
            'end_c_surface_unrounded': float(end[-1]),
            'end_c_mean_unrounded': float(2*w @ end),
            'end_radius_cm': float(self.model.inputs.radius(
                self.end_time, self.model.config.moving))*100,
            'c_min': float(c.min()), 'c_max': float(c.max()),
            't_min': float(temp.min()), 't_max': float(temp.max()),
            'max_outward_c_increase': float(np.max(np.diff(c, axis=0))),
            'max_accumulated_balance_residual': float(abs(invariant).max()),
            'max_instantaneous_balance_residual': float(max(residuals)),
            'radius_extrapolated': bool(self.model.config.moving and
                                        self.end_time > 259200),
            'elapsed_seconds': self.elapsed,
            'nfev': sum(x.nfev for x in self.segments),
            'accepted_steps': sum(len(x.t)-1 for x in self.segments),
        }


def integrate(model, duration=14*86400., stop_at_threshold=True):
    started = time.perf_counter()
    config, m = model.config, model.np
    initial = model.initial_state()
    atol = np.r_[np.full(m, config.atol_t), np.full(m, config.atol_c), config.atol_c]
    def event(t, state):
        return np.max(state[m:2*m]) - .15
    event.terminal = True
    event.direction = -1
    segments, event_time = [], None
    current = 0.
    for end in sorted(set([min(14400., duration), float(duration)])):
        if end <= current:
            continue
        segment = solve_ivp(
            model.rhs, (current, end), initial, method='BDF', dense_output=True,
            events=event if stop_at_threshold else None,
            rtol=config.rtol, atol=atol, jac=model.jacobian,
            max_step=config.early_max_step if end <= 14400 else config.late_max_step)
        if not segment.success:
            raise RuntimeError(segment.message)
        segments.append(segment)
        current, initial = float(segment.t[-1]), segment.y[:, -1]
        if stop_at_threshold and len(segment.t_events[0]):
            event_time = float(segment.t_events[0][0])
            operational = math.floor(event_time) + 1
            final = solve_ivp(model.rhs, (event_time, operational), initial,
                              method='BDF', dense_output=True, rtol=config.rtol,
                              atol=atol, jac=model.jacobian, max_step=.2)
            if not final.success or not np.max(final.y[m:2*m, -1]) < .15:
                raise RuntimeError('Strict endpoint not verified')
            segments.append(final)
            current = float(operational)
            break
    if stop_at_threshold and event_time is None:
        raise RuntimeError(f'Threshold not reached in {duration/3600:g} hours')
    result = Trajectory(model, segments, event_time, current,
                        time.perf_counter()-started)
    check = result.diagnostics()
    if (check['c_min'] <= 0 or check['c_max'] > 2.55001
            or check['t_min'] < 27.9999 or check['t_max'] > 50.2461
            or check['max_accumulated_balance_residual'] > 1e-7):
        raise RuntimeError(f'Physical/conservation validation failed: {check}')
    return result
