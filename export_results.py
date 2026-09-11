"""Expand the official result templates with chunked, four-decimal numerical data."""
import json
import math
import re
import zipfile
from pathlib import Path
import numpy as np
import xlsxwriter
from openpyxl import load_workbook
from drying_model import Inputs, Config, RadialModel, integrate


def compact_dense_workbook(filename):
    """Omit optional cell coordinates in two fully populated dense worksheets.

    OOXML infers the next column when c@r is absent. Values, types, row numbers,
    dimensions and styles are unchanged; only this dense result2 layout qualifies.
    """
    filename = Path(filename)
    temporary = filename.with_name(filename.stem+'-compact.tmp.xlsx')
    with zipfile.ZipFile(filename) as source, zipfile.ZipFile(
            temporary,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:
                rows = re.findall(rb'<row[^>]*>(.*?)</row>',payload)
                for row in rows:
                    coordinates = re.findall(rb'<c r="([A-Z]+)[0-9]+"',row)
                    if coordinates != [bytes([k]) for k in range(65,87)]:
                        raise ValueError('Cannot infer coordinates in a sparse row')
                payload = re.sub(rb'(<c) r="[A-Z]+[0-9]+"',rb'\1',payload)
            target.writestr(item.filename,payload)
    temporary.replace(filename)


def expand_template(template_path, output, trajectory, interval, fields, moving=False):
    template = load_workbook(template_path, read_only=False, data_only=False)
    names = template.sheetnames
    if any(sheet._charts or sheet._images or sheet.merged_cells.ranges for sheet in template):
        raise ValueError('Template requires a more specific preservation route')
    for sheet in template:
        if any(c.data_type == 'f' for row in sheet for c in row):
            raise ValueError('Unexpected formula in result template')
    if list(fields) != names:
        raise ValueError(f'Worksheet names differ from official template: {names}')
    ntime = math.floor(trajectory.end_time/interval)
    times = np.arange(0, ntime+1, dtype=float)*interval
    if times[-1] < trajectory.end_time:
        times = np.r_[times, trajectory.end_time]
    if len(times)+1 > 1048576:
        raise ValueError('Required trajectory exceeds one Excel worksheet')
    book = xlsxwriter.Workbook(output, {'constant_memory': True})
    formats = {'header': book.add_format({'bold': True, 'align': 'center',
                    'valign': 'vcenter', 'text_wrap': True, 'bottom': 1,
                    'font_name': template.worksheets[0]['A1'].font.name or '宋体'}),
               'number': book.add_format({'num_format': '0.0000', 'align': 'center'}),
               'time': book.add_format({'num_format': '0.####', 'align': 'center'})}
    radii = np.arange(21, dtype=float)/10
    headers = [template.worksheets[0]['A1'].value, *radii.tolist()]
    if moving:
        headers.append('药材表面')
    for name, field in fields.items():
        sheet = book.add_worksheet(name)
        sheet.freeze_panes(1, 1)
        sheet.set_column(0, 0, 24)
        sheet.set_column(1, len(headers)-1, 11)
        sheet.set_row(0, 30)
        sheet.write_row(0, 0, headers, formats['header'])
        row_number = 1
        for start in range(0, len(times), 1500):
            chunk = times[start:start+1500]
            temp, moisture = trajectory.sample(chunk, radii)
            values = temp if field == 'temperature' else moisture
            states = trajectory.evaluate(chunk)
            surface_index = (trajectory.model.np-1 if field == 'temperature'
                             else 2*trajectory.model.np-1)
            for i, seconds in enumerate(chunk):
                sheet.write_number(row_number, 0, seconds, formats['time'])
                row = [round(float(v), 4) if np.isfinite(v) else None for v in values[i]]
                if moving:
                    row.append(round(float(states[surface_index, i]), 4))
                sheet.write_row(row_number, 1, row, formats['number'])
                row_number += 1
    notes = book.add_worksheet('说明')
    notes.set_column(0, 0, 28)
    notes.set_column(1, 1, 92)
    text = [
        ('时间单位', '秒，包含初值0；等间隔输出后追加严格达标整秒终点'),
        ('径向位置单位', '厘米，求解器内部使用米'),
        ('温度单位', '摄氏度；经验扩散公式内部换算开尔文'),
        ('含水率单位', 'kg水/kg干物质（干基）'),
        ('输出精度', '数值按题意四舍五入至小数点后四位；不表示物理误差为0.0001'),
        ('环境观测与外推', '前14400秒线性插值；之后固定为末1小时温度和浓度均值'),
        ('网格区间数', trajectory.model.config.n),
        ('物性方案', trajectory.model.config.law),
        ('半径变化', '附件2的保形插值' if moving else '固定2厘米'),
        ('严格不等号', '终点按未舍入的全域最大含水率判定，小数表可能显示为0.1500'),
    ]
    if trajectory.event_time is not None:
        d = trajectory.diagnostics()
        text.extend([('阈值穿越时间/s', trajectory.event_time),
                     ('严格达标时间/s', trajectory.end_time),
                     ('终点最大含水率未舍入', d['end_c_max_unrounded'])])
    if moving:
        text.append(('材料外空间', '固定物理半径超过当前半径的单元留空；末列是当前表面'))
        surface = book.add_worksheet('表面半径')
        surface.write_row(0, 0, ['时间/s', '药材表面半径/cm'], formats['header'])
        surface.set_column(0, 1, 20)
        for i, (t, r) in enumerate(zip(times, trajectory.model.inputs.radius(times, True)*100), 1):
            surface.write_number(i, 0, t, formats['time'])
            surface.write_number(i, 1, float(r), formats['number'])
    for row, values in enumerate(text):
        notes.write_row(row, 0, values)
    book.close()
    template.close()
    if Path(output).name == 'result2.xlsx':
        compact_dense_workbook(output)
    return {'file': Path(output).name, 'sheets': names, 'data_rows': len(times),
            'physical_columns': len(radii), 'surface_column': moving,
            'first_time': float(times[0]), 'last_time': float(times[-1]),
            'interval_seconds': interval, 'bytes': Path(output).stat().st_size}


def export_all(input_dir, output_root):
    output_root = Path(output_root)
    out = output_root/'results'
    summary = json.loads((out/'summary.json').read_text())
    data = Inputs(input_dir)
    receipt = []
    for law, jobs in [('q1', [(1, 1)]), ('q23', [(2, 1), (3, 60)]), ('q4', [(4, 60)])]:
        cfg = Config(**summary['nominal'][law]['config'])
        trajectory = integrate(RadialModel(data, cfg),
                                duration=1800 if law == 'q1' else 14*86400,
                                stop_at_threshold=law != 'q1')
        if law != 'q1':
            assert abs(trajectory.event_time-summary['nominal'][law]['threshold_seconds']) < .01
        for number, interval in jobs:
            fields = ({'温度': 'temperature', '水分浓度': 'moisture'} if number <= 2
                      else {'Sheet1': 'moisture'})
            record = expand_template(Path(input_dir)/'附件3'/f'result{number}.xlsx',
                                      out/f'result{number}.xlsx', trajectory, interval,
                                      fields, moving=number == 4)
            receipt.append(record)
            print(json.dumps(record, ensure_ascii=False), flush=True)
    (out/'excel-export.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2))
    return receipt


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--input-dir', type=Path, default=Path('data/附件'))
    parser.add_argument('--out', type=Path, default=Path('.'))
    args = parser.parse_args()
    export_all(args.input_dir, args.out)
